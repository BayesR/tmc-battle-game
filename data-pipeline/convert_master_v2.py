"""
TMC NPC対戦用カードプール 変換スクリプト（v2：tmc-battle-gameRarity対応）
=========================================
convert_master.py をベースに、独自レアリティ列（tmc-battle-gameRarity）を
NPC_CardPool.xlsx / cardPool.json に引き継ぐよう拡張したもの。
"""

import json
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

SRC_CSV = "/mnt/user-data/uploads/TMC_Database_-_10_CardMaster-4.csv"
OUT_XLSX = "/home/claude/battlestreet/TMC_NPC_CardPool_v2.xlsx"
OUT_JSON = "/home/claude/battlestreet/cardPool_v2.json"

df = pd.read_csv(SRC_CSV, header=1)
df = df.drop(df.columns[0], axis=1)

real = df[df['REGACY'].notna() & (df['REGACY'] != 'REGACY（愛・制・環・聖・邪）')].copy()

for c in ['PP1_Regacy','PP1_Value','PP2_Regacy','PP2_Value','PP3_Regacy','PP3_Value']:
    real[c] = real[c].fillna('')

rarity_rank = {'N': 0, 'R': 1, 'SR': 2, 'UR': 3}
real['rarity_rank'] = real['Rarity'].map(rarity_rank).fillna(0)
real_sorted = real.sort_values('rarity_rank')

key_cols = ['MonsterName','REGACY','MonsterPride','PP1_Regacy','PP1_Value','PP2_Regacy','PP2_Value','PP3_Regacy','PP3_Value','Void']
dedup = real_sorted.groupby(key_cols, dropna=False, as_index=False).agg({
    'CardID': 'first',
    'Rarity': 'first',
    'Product': 'first',
    'CardNo': 'first',
    'tmc-battle-gameRarity': 'first',
})
counts = real_sorted.groupby(key_cols, dropna=False).size().reset_index(name='ReprintCount')
dedup = dedup.merge(counts, on=key_cols, how='left')

rarity_to_level = {'N': 'Lv1', 'R': 'Lv2', 'SR': 'Lv3', 'UR': 'Lv4'}
dedup['SuggestedNPCLevel'] = dedup['Rarity'].map(rarity_to_level).fillna('Lv1')

dedup = dedup.sort_values('MonsterName').reset_index(drop=True)
dedup = dedup[['MonsterName','REGACY','MonsterPride','PP1_Regacy','PP1_Value','PP2_Regacy','PP2_Value',
               'PP3_Regacy','PP3_Value','Void','Rarity','tmc-battle-gameRarity','SuggestedNPCLevel',
               'CardID','Product','ReprintCount']]

print("Final unique cards:", len(dedup))
print(dedup['tmc-battle-gameRarity'].value_counts())

# ---------------------------------------------------------------
# Build workbook (NPC_CardPool, 従来形式＋独自レアリティ列)
# ---------------------------------------------------------------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "NPC_CardPool"

FONT_NAME = "Arial"
header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
header_fill = PatternFill("solid", fgColor="4B5563")
normal_font = Font(name=FONT_NAME, size=10)
battle_fill = PatternFill("solid", fgColor="E0E7FF")  # 独自レアリティ列を色分け表示
thin = Side(style="thin", color="D1D5DB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

headers = ["カード名","Legacy","MonsterPride","PP1_Legacy","PP1_Value","PP2_Legacy","PP2_Value",
           "PP3_Legacy","PP3_Value","Void","本家レアリティ","バトルストリート用レアリティ",
           "使用NPCレベル（提案）","元CardID","収録弾","再録数"]

for col_idx, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border

for r_idx, row in enumerate(dedup.itertuples(index=False), start=2):
    values = list(row)
    for c_idx, val in enumerate(values, start=1):
        cell = ws.cell(row=r_idx, column=c_idx, value=val)
        cell.font = normal_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center" if c_idx not in (1,) else "left")
        if c_idx == 12:
            cell.fill = battle_fill

widths = [20, 8, 12, 10, 9, 10, 9, 10, 9, 8, 11, 16, 12, 10, 26, 8]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A2"

last_row = len(dedup) + 1
def add_validation(formula, col_letter):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{last_row}")

add_validation('"環,愛,制,邪,聖"', "B")
add_validation('"未使用,環,愛,制,邪,聖"', "D")
add_validation('"未使用,環,愛,制,邪,聖"', "F")
add_validation('"未使用,環,愛,制,邪,聖"', "H")
add_validation('"TRUE,FALSE"', "J")
add_validation('"N,R,SR,UR"', "K")
add_validation('"N,R,SR,UR,SUR,SSUR,SSSUR"', "L")
add_validation('"Lv1,Lv2,Lv3,Lv4,Lv5"', "M")

wb.save(OUT_XLSX)
print("saved xlsx:", OUT_XLSX)

# ---------------------------------------------------------------
# Build cardPool.json（アプリ用。battleStreetRarity フィールドを追加）
# ---------------------------------------------------------------
def to_num(v):
    if v is None or v == '':
        return 0
    try:
        return float(v)
    except Exception:
        return 0

cards = []
used_ids = set()
for _, r in dedup.iterrows():
    def slot(l, v):
        return {'legacy': l if l else '未使用', 'value': to_num(v)}
    pps = [slot(r['PP1_Regacy'], r['PP1_Value']), slot(r['PP2_Regacy'], r['PP2_Value']), slot(r['PP3_Regacy'], r['PP3_Value'])]
    cid = str(r['CardID']) if pd.notna(r['CardID']) else f'GEN_{len(cards):04d}'
    base_id = cid
    i = 1
    while cid in used_ids:
        cid = f'{base_id}_{i}'
        i += 1
    used_ids.add(cid)
    cards.append({
        'id': cid,
        'name': r['MonsterName'],
        'legacy': r['REGACY'],
        'monsterPride': int(to_num(r['MonsterPride'])),
        'potentialPoints': pps,
        'hasVoid': (str(r['Void']).strip().upper() == 'TRUE'),
        'rarity': r['Rarity'] if pd.notna(r['Rarity']) else 'N',
        'suggestedNpcLevel': r['SuggestedNPCLevel'],
        'battleStreetRarity': r['tmc-battle-gameRarity'],
    })

with open(OUT_JSON, 'w', encoding='utf-8') as f:
    json.dump(cards, f, ensure_ascii=False, indent=2)
print("saved json:", OUT_JSON, "count:", len(cards))
